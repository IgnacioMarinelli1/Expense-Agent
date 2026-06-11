import io
from datetime import datetime
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import APIRouter, Query
from fastapi.responses import Response

from db.db import get_db
from db.security import current_user_id

router = APIRouter(prefix="/export", tags=["export"])


def _notes_from_doc(doc: dict) -> str:
    notes = doc.get("notes")
    if notes:
        return notes
    metadata_notes = (doc.get("metadata") or {}).get("notas")
    if metadata_notes:
        return metadata_notes
    currency = doc.get("currency", "ARS")
    return f"Gasto {currency}" if currency != "ARS" else "Gasto sin descripción"


def _fmt_date(dt) -> str:
    if not dt:
        return ""
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt)
        except ValueError:
            return dt
    return dt.strftime("%d/%m/%Y")


@router.get("/expenses")
async def export_expenses(month: Optional[str] = Query(None, description="YYYY-MM")):
    db = get_db()
    user_id = current_user_id()

    query: dict = {"user_id": user_id, "type": {"$ne": "income"}}
    if month:
        query["period"] = month

    cursor = db["payments"].find(query).sort("payment_date", -1)
    docs = await cursor.to_list(length=2000)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gastos"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    headers = ["Fecha", "Descripción", "Monto", "Moneda", "Estado", "Período"]
    col_widths = [14, 40, 16, 10, 12, 10]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 20

    status_map = {"paid": "Pagado", "pending": "Pendiente", "overdue": "Vencido"}

    for row_idx, doc in enumerate(docs, 2):
        ws.cell(row=row_idx, column=1, value=_fmt_date(doc.get("payment_date")))
        ws.cell(row=row_idx, column=2, value=_notes_from_doc(doc))
        amount_cell = ws.cell(row=row_idx, column=3, value=doc.get("amount", 0))
        amount_cell.number_format = '#,##0.00'
        ws.cell(row=row_idx, column=4, value=doc.get("currency", "ARS"))
        ws.cell(row=row_idx, column=5, value=status_map.get(doc.get("status", ""), doc.get("status", "")))
        ws.cell(row=row_idx, column=6, value=doc.get("period", ""))

    if docs:
        total_row = len(docs) + 2
        ws.cell(row=total_row, column=2, value="TOTAL").font = Font(bold=True)
        total_cell = ws.cell(row=total_row, column=3, value=sum(d.get("amount", 0) for d in docs))
        total_cell.font = Font(bold=True)
        total_cell.number_format = '#,##0.00'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    period_label = month or "todos"
    filename = f"gastos_{period_label}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
